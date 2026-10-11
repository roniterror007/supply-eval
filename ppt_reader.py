from pptx import Presentation

prs = Presentation('Agentic_AI_Demand_to_Supply_Hackathon_Proposal.pptx')
text_runs = []
for i, slide in enumerate(prs.slides):
    text_runs.append(f'Slide {i}:')
    for shape in slide.shapes:
        if hasattr(shape, 'text'):
            text_runs.append(shape.text)

with open('ppt_text.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(text_runs))
